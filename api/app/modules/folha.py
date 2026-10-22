"""Exportação do consumo para a importação de eventos da folha do Sankhya.

Os nomes, a ordem e o TIPO de cada coluna não são escolha nossa: são o
contrato do ERP, que lê a tipagem da célula na importação. Um número gravado
como texto, ou uma data gravada como string, é recusado lá. Por isso o tipo
está declarado em `COLUNAS` e não deixado por conta do acaso.

Cada pessoa gera até duas linhas por competência — uma por evento — porque a
folha separa o que foi consumido na loja do que foi consumido no refeitório.
Quem não consumiu de um lado não gera a linha daquele lado.
"""

import uuid
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select
from sqlalchemy.orm import Session, aliased

from app.core.config import obter_config
from app.excecoes import SemPermissao
from app.models.cadastro import Colaborador, Departamento, Empresa
from app.models.operacao import Almoco, ItemPedido, Pedido
from app.modules import ciclo
from app.modules.auditoria import Ator

# Eventos da folha. São dois porque o holerite mostra as duas coisas separadas.
CODEVENTO_LOJA = 10206  # CONSUMOS DIVERSOS
CODEVENTO_REFEITORIO = 10221  # CONSUMO REFEITORIO

# Constantes do layout: valem para toda linha, sempre.
TIPMOV = "M"
SEQUENCIA = 0
INDICE = Decimal("0.00")
TIPEVENTO = -1 
UNIDADE = "V"
TEXTO = "@"
INTEIRO = "0"
DECIMAL = "0.00"
DATA = "DD/MM/YYYY"

# (nome da coluna, formato, largura)
COLUNAS: tuple[tuple[str, str, int], ...] = (
    ("CODEMP", INTEIRO, 9),
    ("CODFUNC", INTEIRO, 10),
    ("REFERENCIA", DATA, 13),
    ("TIPMOV", TEXTO, 9),
    ("CODEVENTO", INTEIRO, 12),
    ("SEQUENCIA", INTEIRO, 12),
    ("INDICE", DECIMAL, 9),
    ("VLRMOV", DECIMAL, 12),
    ("TIPEVENTO", INTEIRO, 12),
    ("UNIDADE", TEXTO, 10),
)


@dataclass
class LinhaFolha:
    """Uma linha do arquivo. CODFUNC é a matrícula gravada no movimento."""

    codemp: int
    codfunc: int
    codevento: int
    valor: Decimal


@dataclass
class LinhaSemMatricula:
    """Quem consumiu e não cabe no arquivo da folha, por não ter matrícula.

    O PJ é proibido de ter matrícula por regra do cadastro, então ele nunca
    entra na importação de eventos. Sai numa planilha à parte para ser cobrado
    por fora — o consumo existe do mesmo jeito e não pode evaporar.
    """

    codemp: int
    codparc: int
    nome: str
    total_loja: Decimal
    total_refeitorio: Decimal

    @property
    def total(self) -> Decimal:
        return self.total_loja + self.total_refeitorio


@dataclass
class LinhaCortesiaIsento:
    """Almoço que ninguém paga do próprio bolso: visitante liberado de graça,
    ou funcionário isento (a empresa cobre, não desconta da pessoa).

    Existe pro DP ver quanto a empresa está bancando no refeitório — nenhuma
    das duas entra na soma por pessoa nem no arquivo da folha.
    """

    tipo: str
    nome: str
    data: date
    valor: Decimal


@dataclass
class ConsumoDaPessoa:
    """O que uma pessoa consumiu no ciclo, antes de virar linha de arquivo.

    É desta lista que saem as três visões: as linhas da folha, quem ficou de
    fora por não ter matrícula, e a conferência que o DP usa para responder
    "por que descontaram isso de mim" sem abrir a planilha.
    """

    codemp: int
    codparc: int
    matricula: int | None
    nome: str
    codigo: str
    loja: Decimal
    refeitorio: Decimal

    @property
    def total(self) -> Decimal:
        return self.loja + self.refeitorio


@dataclass
class Folha:
    competencia: str
    referencia: date
    pessoas: list[ConsumoDaPessoa]
    cortesias_isentos: list[LinhaCortesiaIsento]

    @property
    def linhas(self) -> list[LinhaFolha]:
        """As linhas do arquivo. Valor zerado não vira linha: um desconto de
        R$ 0,00 no holerite é ruído."""
        linhas: list[LinhaFolha] = []
        for pessoa in self.pessoas:
            if pessoa.matricula is None:
                continue
            if pessoa.loja:
                linhas.append(
                    LinhaFolha(pessoa.codemp, pessoa.matricula, CODEVENTO_LOJA, pessoa.loja)
                )
            if pessoa.refeitorio:
                linhas.append(
                    LinhaFolha(
                        pessoa.codemp, pessoa.matricula, CODEVENTO_REFEITORIO, pessoa.refeitorio
                    )
                )
        return linhas

    @property
    def sem_matricula(self) -> list[LinhaSemMatricula]:
        return [
            LinhaSemMatricula(
                codemp=pessoa.codemp,
                codparc=pessoa.codparc,
                nome=pessoa.nome,
                total_loja=pessoa.loja,
                total_refeitorio=pessoa.refeitorio,
            )
            for pessoa in self.pessoas
            if pessoa.matricula is None
        ]

    @property
    def total(self) -> Decimal:
        return sum((linha.valor for linha in self.linhas), Decimal("0.00"))

    @property
    def total_cortesias_isentos(self) -> Decimal:
        return sum((linha.valor for linha in self.cortesias_isentos), Decimal("0.00"))


def _referencia(competencia: str) -> date:
    """Primeiro dia do mês em que o ciclo fecha.

    A competência 2026-08 cobre de 21/07 a 20/08 e é referenciada na folha
    como 01/08/2026 — o mês, não o intervalo.
    """
    ano, mes = (int(parte) for parte in competencia.split("-"))
    return date(ano, mes, 1)

def _cortesias_e_isentos(
    sessao: Session, competencia: str
) -> list[LinhaCortesiaIsento]:
    visitantes = sessao.execute(
        select(Almoco.visitante_nome, Almoco.criado_em, Almoco.valor).where(
            Almoco.cortesia.is_(True),
            Almoco.status == "confirmado",
            Almoco.competencia == competencia,
        )
    ).all()
    isentos = sessao.execute(
        select(Colaborador.nome_completo, Almoco.criado_em, Almoco.valor)
        .join(Colaborador, Colaborador.id == Almoco.colaborador_id)
        .where(
            Almoco.isento.is_(True),
            Almoco.status == "confirmado",
            Almoco.competencia == competencia,
        )
    ).all()
    return [
        LinhaCortesiaIsento("visitante", nome, criado_em.date(), valor)
        for nome, criado_em, valor in visitantes
    ] + [
        LinhaCortesiaIsento("isento", nome, criado_em.date(), valor)
        for nome, criado_em, valor in isentos
    ]


type _Chave = tuple[uuid.UUID, uuid.UUID, int | None]


def _somar(sessao: Session, modelo, coluna_valor, condicoes) -> dict[_Chave, Decimal]:
    linhas = sessao.execute(
        select(
            modelo.colaborador_id,
            modelo.empresa_id,
            modelo.matricula,
            func.sum(coluna_valor),
        )
        .where(*condicoes)
        .group_by(modelo.colaborador_id, modelo.empresa_id, modelo.matricula)
    ).all()
    return {(pessoa, empresa, matricula): valor for pessoa, empresa, matricula, valor in linhas}


def montar(
    sessao: Session, ator: Ator, competencia: str, lote_id: uuid.UUID | None = None
) -> Folha:
    """Consolida o ciclo em linhas de evento, uma por pessoa e por evento.

    As regras de quem conta são as mesmas do extrato que o colaborador vê:
    pedido não cancelado e almoço confirmado. Precisa ser a mesma, senão a
    pessoa confere o próprio consumo na tela e encontra outro número no
    holerite.
    """
    if ator.papel not in {"admin", "dp"}:
        raise SemPermissao()


    ciclo.validar(competencia)

    # Visitante não tem colaborador_id: não tem folha pra descontar, então
    # nunca pode entrar nesta soma. Sem este filtro, `pessoas[colaborador_id]`
    # explode com KeyError assim que um visitante aparecer na competência.
    if lote_id is not None:
        onde_loja = [Pedido.lote_id == lote_id, Pedido.colaborador_id.isnot(None)]
        onde_refeitorio = [
            Almoco.lote_id == lote_id,
            Almoco.colaborador_id.isnot(None),
            Almoco.isento.is_(False),
        ]
    else:
        onde_loja = [
            Pedido.competencia == competencia,
            Pedido.status != "cancelado",
            Pedido.colaborador_id.isnot(None),
        ]
        onde_refeitorio = [
            Almoco.competencia == competencia,
            Almoco.status == "confirmado",
            Almoco.colaborador_id.isnot(None),
            # Isento: a empresa paga, não desconta da pessoa — vira linha
            # separada em `cortesias_isentos`, não entra nesta soma.
            Almoco.isento.is_(False),
        ]

    loja = _somar(sessao, Pedido, Pedido.valor_total, onde_loja)
    refeitorio = _somar(sessao, Almoco, Almoco.valor, onde_refeitorio)

    chaves = set(loja) | set(refeitorio)
    codemp_de = dict(sessao.execute(select(Empresa.id, Empresa.codemp)).all())
    pessoas = {
        pessoa.id: pessoa
        for pessoa in sessao.scalars(
            select(Colaborador).where(Colaborador.id.in_({chave[0] for chave in chaves}))
        )
    }

    consolidado: list[ConsumoDaPessoa] = []
    for chave in chaves:
        colaborador_id, empresa_id, matricula = chave
        pessoa = pessoas[colaborador_id]
        consolidado.append(
            ConsumoDaPessoa(
                codemp=codemp_de[empresa_id],
                codparc=pessoa.codparc,
                matricula=matricula,
                nome=pessoa.nome_completo,
                codigo=pessoa.codigo,
                loja=loja.get(chave, Decimal("0.00")),
                refeitorio=refeitorio.get(chave, Decimal("0.00")),
            )
        )

    return Folha(
        competencia=competencia,
        referencia=_referencia(competencia),
        # Com matrícula primeiro e em ordem de CODFUNC: é assim que a folha lê,
        # e é a ordem em que o DP confere linha a linha contra o ERP.
        pessoas=sorted(
            consolidado, key=lambda p: (p.matricula is None, p.matricula or 0, p.nome)
        ),
        cortesias_isentos=_cortesias_e_isentos(sessao, competencia),
    )


def _escrever(
    titulo: str, colunas: tuple[tuple[str, str, int], ...], fileiras: list[list]
) -> bytes:
    """Monta a planilha aplicando o formato declarado de cada coluna.

    O formato é aplicado célula a célula, e não à coluna inteira, porque é a
    célula que carrega o tipo para quem lê o arquivo depois.
    """
    livro = Workbook()
    _preencher(livro.active, titulo, colunas, fileiras)

    buffer = BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


def _preencher(
    aba: Worksheet,
    titulo: str,
    colunas: tuple[tuple[str, str, int], ...],
    fileiras: list[list],
) -> None:
    aba.title = titulo

    aba.append([nome for nome, _formato, _largura in colunas])
    for celula in aba[1]:
        celula.font = Font(bold=True)
    aba.freeze_panes = "A2"

    for fileira in fileiras:
        if len(fileira) != len(colunas):
            raise ValueError(f"linha com {len(fileira)} valores para {len(colunas)} colunas")
        aba.append(fileira)

    for indice, (_nome, formato, largura) in enumerate(colunas, start=1):
        aba.column_dimensions[aba.cell(row=1, column=indice).column_letter].width = largura
        for fileira_excel in range(2, aba.max_row + 1):
            aba.cell(row=fileira_excel, column=indice).number_format = formato


def planilha(folha: Folha) -> bytes:
    """O arquivo que sobe no Sankhya. Uma aba só, nada além das colunas dele."""
    return _escrever(
        "FOLHA",
        COLUNAS,
        [
            [
                linha.codemp,
                linha.codfunc,
                folha.referencia,
                TIPMOV,
                linha.codevento,
                SEQUENCIA,
                INDICE,
                linha.valor,
                TIPEVENTO,
                UNIDADE,
            ]
            for linha in folha.linhas
        ],
    )


COLUNAS_SEM_MATRICULA: tuple[tuple[str, str, int], ...] = (
    ("CODEMP", INTEIRO, 9),
    ("CODPARC", INTEIRO, 10),
    ("NOME", TEXTO, 34),
    ("REFERENCIA", DATA, 13),
    ("CONSUMOS DIVERSOS", DECIMAL, 20),
    ("CONSUMO REFEITORIO", DECIMAL, 20),
    ("TOTAL", DECIMAL, 12),
)


def planilha_sem_matricula(folha: Folha) -> bytes:
    """Consumo dos PJ. Layout nosso, não do Sankhya: este arquivo não é
    importado, é lido por gente para cobrar por outro caminho."""
    return _escrever(
        "PJ",
        COLUNAS_SEM_MATRICULA,
        [
            [
                linha.codemp,
                linha.codparc,
                linha.nome,
                folha.referencia,
                linha.total_loja,
                linha.total_refeitorio,
                linha.total,
            ]
            for linha in folha.sem_matricula
        ],
    )


@dataclass
class DiaDeConsumo:
    data: date
    almoco: Decimal
    loja: Decimal
    itens: str
    # Almoço que a empresa paga: aparece na folha, mas não entra no desconto.
    almoco_isento: Decimal = Decimal("0.00")


@dataclass
class FolhaDeAssinatura:
    colaborador_id: uuid.UUID
    departamento_id: uuid.UUID | None
    nome: str
    codigo: str
    matricula: int | None
    empresa: str
    departamento: str | None
    dias: list[DiaDeConsumo]

    @property
    def total(self) -> Decimal:
        """O que vai ser descontado em folha. Almoço isento fica de fora."""
        return sum((dia.almoco + dia.loja for dia in self.dias), Decimal("0.00"))

    @property
    def total_isento(self) -> Decimal:
        return sum((dia.almoco_isento for dia in self.dias), Decimal("0.00"))


def folhas_de_assinatura(
    sessao: Session,
    ator: Ator,
    competencia: str,
    *,
    departamento_id: uuid.UUID | None = None,
    colaborador_id: uuid.UUID | None = None,
) -> list[FolhaDeAssinatura]:
    """Uma folha por pessoa com o consumo dia a dia, para ela conferir e
    assinar. O desconto segue as regras de `montar`. O isento também assina:
    o almoço dele aparece pelo valor real, mas com desconto zero."""
    if ator.papel not in {"admin", "dp"}:
        raise SemPermissao()
    ciclo.validar(competencia)

    fuso = ZoneInfo(obter_config().fuso)
    dias: dict[uuid.UUID, dict[date, DiaDeConsumo]] = {}

    def dia_de(pessoa: uuid.UUID, instante: datetime) -> DiaDeConsumo:
        data = instante.astimezone(fuso).date()
        por_dia = dias.setdefault(pessoa, {})
        return por_dia.setdefault(data, DiaDeConsumo(data, Decimal("0.00"), Decimal("0.00"), ""))

    pedidos = list(
        sessao.scalars(
            select(Pedido).where(
                Pedido.competencia == competencia,
                Pedido.status != "cancelado",
                Pedido.colaborador_id.isnot(None),
            )
        )
    )
    itens: dict[uuid.UUID, list[str]] = {}
    if pedidos:
        for item in sessao.scalars(
            select(ItemPedido).where(ItemPedido.pedido_id.in_([p.id for p in pedidos]))
        ):
            itens.setdefault(item.pedido_id, []).append(f"{item.quantidade}x {item.nome_produto}")
    for pedido in pedidos:
        dia = dia_de(pedido.colaborador_id, pedido.criado_em)
        dia.loja += pedido.valor_total
        dia.itens = ", ".join(filter(None, [dia.itens, *itens.get(pedido.id, [])]))

    for almoco in sessao.scalars(
        select(Almoco).where(
            Almoco.competencia == competencia,
            Almoco.status == "confirmado",
            Almoco.colaborador_id.isnot(None),
        )
    ):
        dia = dia_de(almoco.colaborador_id, almoco.criado_em)
        if almoco.isento:
            dia.almoco_isento += almoco.valor
        else:
            dia.almoco += almoco.valor

    if not dias:
        return []
    pessoas = sessao.execute(
        select(Colaborador, Empresa.nome, Departamento.nome)
        .join(Empresa, Empresa.id == Colaborador.empresa_id)
        .outerjoin(Departamento, Departamento.id == Colaborador.departamento_id)
        .where(Colaborador.id.in_(list(dias)))
    ).all()
    folhas = [
        FolhaDeAssinatura(
            colaborador_id=pessoa.id,
            departamento_id=pessoa.departamento_id,
            nome=pessoa.nome_completo,
            codigo=pessoa.codigo,
            matricula=pessoa.matricula,
            empresa=empresa,
            departamento=departamento,
            dias=sorted(dias[pessoa.id].values(), key=lambda d: d.data),
        )
        for pessoa, empresa, departamento in pessoas
        if (departamento_id is None or pessoa.departamento_id == departamento_id)
        and (colaborador_id is None or pessoa.id == colaborador_id)
    ]
    return sorted(folhas, key=lambda f: ((f.departamento or "").lower(), f.nome.lower()))


@dataclass
class AlmocoSemDesconto:
    data: date
    nome: str
    tipo: str
    departamento: str | None
    pagamento: str
    valor: Decimal
    pago_pela_empresa: Decimal
    pago_pelo_consumidor: Decimal


def _pagamento(almoco: Almoco) -> str:
    if almoco.isento:
        return "Isento (empresa paga)"
    if almoco.cortesia:
        return "Cortesia"
    return "Pix confirmado" if almoco.pix_confirmado_em else "Pix pendente"


def almocos_sem_desconto(
    sessao: Session, ator: Ator, competencia: str
) -> list[AlmocoSemDesconto]:
    """Todo almoço do ciclo que não vira desconto em folha: funcionário isento
    e visitante (cortesia ou Pix). Para o DP ver quanto a empresa banca."""
    if ator.papel not in {"admin", "dp"}:
        raise SemPermissao()
    ciclo.validar(competencia)

    fuso = ZoneInfo(obter_config().fuso)
    area_do_visitante = aliased(Departamento)
    linhas = sessao.execute(
        select(
            Almoco,
            Colaborador.nome_completo,
            Departamento.nome,
            area_do_visitante.nome,
        )
        .outerjoin(Colaborador, Colaborador.id == Almoco.colaborador_id)
        .outerjoin(Departamento, Departamento.id == Colaborador.departamento_id)
        .outerjoin(area_do_visitante, area_do_visitante.id == Almoco.visitante_departamento_id)
        .where(
            Almoco.competencia == competencia,
            Almoco.status == "confirmado",
            Almoco.isento.is_(True) | Almoco.visitante_nome.isnot(None),
        )
    ).all()

    zero = Decimal("0.00")
    resultado = []
    for almoco, nome_colaborador, area, area_visitante in linhas:
        empresa_paga = almoco.isento or almoco.cortesia
        resultado.append(
            AlmocoSemDesconto(
                data=almoco.criado_em.astimezone(fuso).date(),
                nome=almoco.visitante_nome or nome_colaborador,
                tipo="Visitante" if almoco.visitante_nome else "Funcionário isento",
                departamento=area_visitante if almoco.visitante_nome else area,
                pagamento=_pagamento(almoco),
                valor=almoco.valor,
                pago_pela_empresa=almoco.valor if empresa_paga else zero,
                pago_pelo_consumidor=(
                    almoco.valor if almoco.pix_confirmado_em and not empresa_paga else zero
                ),
            )
        )
    return sorted(
        resultado, key=lambda a: ((a.departamento or "").lower(), a.nome.lower(), a.data)
    )


COLUNAS_ALMOCOS_DETALHADO: tuple[tuple[str, str, int], ...] = (
    ("DATA", DATA, 13),
    ("NOME", TEXTO, 34),
    ("TIPO", TEXTO, 20),
    ("DEPARTAMENTO", TEXTO, 26),
    ("PAGAMENTO", TEXTO, 22),
    ("VALOR DO ALMOCO", DECIMAL, 17),
    ("PAGO PELA EMPRESA", DECIMAL, 19),
    ("PAGO PELO CONSUMIDOR", DECIMAL, 22),
)

COLUNAS_ALMOCOS_POR_PESSOA: tuple[tuple[str, str, int], ...] = (
    ("NOME", TEXTO, 34),
    ("TIPO", TEXTO, 20),
    ("DEPARTAMENTO", TEXTO, 26),
    ("ALMOCOS", INTEIRO, 10),
    ("VALOR DOS ALMOCOS", DECIMAL, 19),
    ("PAGO PELA EMPRESA", DECIMAL, 19),
    ("PAGO PELO CONSUMIDOR", DECIMAL, 22),
)


def planilha_almocos_sem_desconto(almocos: list[AlmocoSemDesconto]) -> bytes:
    """Duas abas: o detalhe, um almoço por linha, e o total por pessoa. Layout
    nosso, lido por gente: não é importado em lugar nenhum."""
    por_pessoa: dict[tuple[str, str, str | None], list] = {}
    for a in almocos:
        acumulado = por_pessoa.setdefault(
            (a.nome, a.tipo, a.departamento), [0, Decimal("0.00"), Decimal("0.00"), Decimal("0.00")]
        )
        acumulado[0] += 1
        acumulado[1] += a.valor
        acumulado[2] += a.pago_pela_empresa
        acumulado[3] += a.pago_pelo_consumidor

    linhas_pessoa = [
        [nome, tipo, departamento or "Sem departamento", *numeros]
        for (nome, tipo, departamento), numeros in sorted(
            por_pessoa.items(), key=lambda i: ((i[0][2] or "").lower(), i[0][0].lower())
        )
    ]
    linhas_pessoa.append(
        [
            "TOTAL",
            None,
            None,
            sum(p[3] for p in linhas_pessoa),
            sum((p[4] for p in linhas_pessoa), Decimal("0.00")),
            sum((p[5] for p in linhas_pessoa), Decimal("0.00")),
            sum((p[6] for p in linhas_pessoa), Decimal("0.00")),
        ]
    )

    livro = Workbook()
    _preencher(
        livro.active,
        "DETALHADO",
        COLUNAS_ALMOCOS_DETALHADO,
        [
            [
                a.data,
                a.nome,
                a.tipo,
                a.departamento or "Sem departamento",
                a.pagamento,
                a.valor,
                a.pago_pela_empresa,
                a.pago_pelo_consumidor,
            ]
            for a in almocos
        ],
    )
    _preencher(livro.create_sheet(), "POR PESSOA", COLUNAS_ALMOCOS_POR_PESSOA, linhas_pessoa)
    for celula in livro["POR PESSOA"][livro["POR PESSOA"].max_row]:
        celula.font = Font(bold=True)

    buffer = BytesIO()
    livro.save(buffer)
    return buffer.getvalue()
