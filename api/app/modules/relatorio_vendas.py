"""Relatórios de vendas em Excel — sempre em duas abas, uma empresa separada
do resto do grupo.

A empresa separada sai por pedido de negócio, não por regra técnica: o
CODEMP é só um valor de configuração, não um enum do domínio, por isso vive
aqui como constante e não como caso especial em algum outro módulo.
"""

import re
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import obter_config
from app.excecoes import PeriodoInvalido, SemPermissao
from app.models.cadastro import Colaborador, Departamento, Empresa
from app.models.operacao import Almoco, ItemPedido, Pedido
from app.modules.auditoria import Ator

CODEMP_EMPRESA_SEPARADA = 17

COLUNAS_DETALHADO = (
    "Codparc",
    "Colaborador",
    "Departamento",
    "Empresa",
    "Data",
    "Tipo",
    "Item",
    "Categoria",
    "Quantidade",
    "Valor unitário",
    "Valor total",
    "Código do pedido",
    "Brinde de aniversário",
)

COLUNAS_CONSOLIDADO = (
    "Codparc",
    "Colaborador",
    "Departamento",
    "Empresa",
    "Compras",
    "Itens",
    "Almoços",
    "Valor total",
)


@dataclass
class LinhaVenda:
    codemp: int
    empresa_nome: str
    codparc: int
    colaborador_nome: str
    departamento: str | None
    data: datetime
    tipo: str
    item: str
    categoria: str | None
    quantidade: int
    preco_unitario: Decimal
    valor_total: Decimal
    codigo_pedido: str
    brinde: bool


@dataclass
class LinhaConsolidada:
    codemp: int
    empresa_nome: str
    codparc: int
    colaborador_nome: str
    departamento: str | None
    compras: int
    itens: int
    almocos: int
    valor_total: Decimal


def _intervalo(de: date, ate: date) -> tuple[datetime, datetime]:
    if ate < de:
        raise PeriodoInvalido()
    fuso = ZoneInfo(obter_config().fuso)
    inicio = datetime(de.year, de.month, de.day, tzinfo=fuso).astimezone(UTC)
    fim = (datetime(ate.year, ate.month, ate.day, tzinfo=fuso) + timedelta(days=1)).astimezone(UTC)
    return inicio, fim


def montar(sessao: Session, ator: Ator, de: date, ate: date) -> list[LinhaVenda]:
    """Só o que já virou entrega/confirmação de verdade — é dado pra terceiro,
    não pode incluir pedido ainda pendente."""
    if not ator.eh_admin:
        raise SemPermissao()
    inicio, fim = _intervalo(de, ate)

    linhas: list[LinhaVenda] = []

    consulta_loja = (
        select(Pedido, Colaborador, Departamento.nome, Empresa.codemp, Empresa.nome, ItemPedido)
        .join(Colaborador, Colaborador.id == Pedido.colaborador_id)
        .outerjoin(Departamento, Departamento.id == Colaborador.departamento_id)
        .join(Empresa, Empresa.id == Pedido.empresa_id)
        .join(ItemPedido, ItemPedido.pedido_id == Pedido.id)
        .where(Pedido.status == "entregue", Pedido.criado_em >= inicio, Pedido.criado_em < fim)
        .order_by(Colaborador.nome_completo, Pedido.criado_em)
    )
    for pedido, colaborador, departamento, codemp, empresa_nome, item in sessao.execute(
        consulta_loja
    ).all():
        linhas.append(
            LinhaVenda(
                codemp=codemp,
                empresa_nome=empresa_nome,
                codparc=colaborador.codparc,
                colaborador_nome=colaborador.nome_completo,
                departamento=departamento,
                data=pedido.criado_em,
                tipo="Compra na loja",
                item=item.nome_produto,
                categoria=item.categoria,
                quantidade=item.quantidade,
                preco_unitario=item.preco_unitario,
                valor_total=item.preco_unitario * item.quantidade,
                codigo_pedido=pedido.codigo_retirada,
                brinde=item.brinde,
            )
        )

    consulta_refeitorio = (
        select(Almoco, Colaborador, Departamento.nome, Empresa.codemp, Empresa.nome)
        .join(Colaborador, Colaborador.id == Almoco.colaborador_id)
        .outerjoin(Departamento, Departamento.id == Colaborador.departamento_id)
        .join(Empresa, Empresa.id == Almoco.empresa_id)
        .where(Almoco.status == "confirmado", Almoco.criado_em >= inicio, Almoco.criado_em < fim)
        .order_by(Colaborador.nome_completo, Almoco.criado_em)
    )
    for almoco, colaborador, departamento, codemp, empresa_nome in sessao.execute(
        consulta_refeitorio
    ).all():
        linhas.append(
            LinhaVenda(
                codemp=codemp,
                empresa_nome=empresa_nome,
                codparc=colaborador.codparc,
                colaborador_nome=colaborador.nome_completo,
                departamento=departamento,
                data=almoco.criado_em,
                tipo="Almoço",
                item="Almoço no refeitório",
                categoria=None,
                quantidade=1,
                preco_unitario=almoco.valor,
                valor_total=almoco.valor,
                codigo_pedido="—",
                brinde=False,
            )
        )

    return linhas


def consolidar_por_pessoa(linhas: list[LinhaVenda]) -> list[LinhaConsolidada]:
    """Uma linha por pessoa: total gasto no período, loja e refeitório juntos.

    'Compras' conta pedido distinto, não item — uma compra de três produtos é
    uma compra só. Por isso o código do pedido de cada um é guardado num
    conjunto à parte, e não somado direto como itens e almoços são.
    """
    agregados: dict[int, dict] = {}
    pedidos_por_pessoa: dict[int, set[str]] = {}

    for linha in linhas:
        dados = agregados.setdefault(
            linha.codparc,
            {
                "codemp": linha.codemp,
                "empresa_nome": linha.empresa_nome,
                "colaborador_nome": linha.colaborador_nome,
                "departamento": linha.departamento,
                "itens": 0,
                "almocos": 0,
                "valor_total": Decimal("0"),
            },
        )
        dados["valor_total"] += linha.valor_total
        if linha.tipo == "Almoço":
            dados["almocos"] += 1
        else:
            dados["itens"] += linha.quantidade
            pedidos_por_pessoa.setdefault(linha.codparc, set()).add(linha.codigo_pedido)

    return [
        LinhaConsolidada(
            codemp=dados["codemp"],
            empresa_nome=dados["empresa_nome"],
            codparc=codparc,
            colaborador_nome=dados["colaborador_nome"],
            departamento=dados["departamento"],
            compras=len(pedidos_por_pessoa.get(codparc, set())),
            itens=dados["itens"],
            almocos=dados["almocos"],
            valor_total=dados["valor_total"],
        )
        for codparc, dados in agregados.items()
    ]


def _nome_da_empresa_separada(sessao: Session) -> str:
    nome = sessao.scalar(select(Empresa.nome).where(Empresa.codemp == CODEMP_EMPRESA_SEPARADA))
    return nome or f"Empresa {CODEMP_EMPRESA_SEPARADA}"


def _nome_de_aba(bruto: str) -> str:
    """Excel proíbe alguns caracteres no título da aba e corta em 31."""
    limpo = re.sub(r'[:\\/?*\[\]]', "", bruto).strip()
    return (limpo or f"Empresa {CODEMP_EMPRESA_SEPARADA}")[:31]


def _escrever_detalhado(aba: Worksheet, linhas: list[LinhaVenda]) -> None:
    aba.append(list(COLUNAS_DETALHADO))
    for celula in aba[1]:
        celula.font = Font(bold=True)
    aba.freeze_panes = "A2"

    for linha in linhas:
        aba.append(
            [
                linha.codparc,
                linha.colaborador_nome,
                linha.departamento or "—",
                f"{linha.codemp} — {linha.empresa_nome}",
                linha.data.strftime("%d/%m/%Y %H:%M"),
                linha.tipo,
                linha.item,
                linha.categoria or "—",
                linha.quantidade,
                float(linha.preco_unitario),
                float(linha.valor_total),
                linha.codigo_pedido,
                "Sim" if linha.brinde else "—",
            ]
        )

    larguras = (12, 28, 22, 28, 16, 16, 28, 16, 10, 14, 14, 16, 18)
    for indice, largura in enumerate(larguras, start=1):
        aba.column_dimensions[aba.cell(row=1, column=indice).column_letter].width = largura


def _escrever_consolidado(aba: Worksheet, linhas: list[LinhaConsolidada]) -> None:
    aba.append(list(COLUNAS_CONSOLIDADO))
    for celula in aba[1]:
        celula.font = Font(bold=True)
    aba.freeze_panes = "A2"

    for linha in linhas:
        aba.append(
            [
                linha.codparc,
                linha.colaborador_nome,
                linha.departamento or "—",
                f"{linha.codemp} — {linha.empresa_nome}",
                linha.compras,
                linha.itens,
                linha.almocos,
                float(linha.valor_total),
            ]
        )

    larguras = (12, 28, 22, 28, 10, 10, 10, 14)
    for indice, largura in enumerate(larguras, start=1):
        aba.column_dimensions[aba.cell(row=1, column=indice).column_letter].width = largura


def planilha_detalhada(sessao: Session, linhas: list[LinhaVenda]) -> bytes:
    livro = Workbook()
    aba_demais: Worksheet = livro.active
    aba_demais.title = "Demais empresas"
    _escrever_detalhado(
        aba_demais, [linha for linha in linhas if linha.codemp != CODEMP_EMPRESA_SEPARADA]
    )

    aba_separada = livro.create_sheet(_nome_de_aba(_nome_da_empresa_separada(sessao)))
    _escrever_detalhado(
        aba_separada, [linha for linha in linhas if linha.codemp == CODEMP_EMPRESA_SEPARADA]
    )

    buffer = BytesIO()
    livro.save(buffer)
    return buffer.getvalue()


def planilha_consolidada(sessao: Session, linhas: list[LinhaConsolidada]) -> bytes:
    livro = Workbook()
    aba_demais: Worksheet = livro.active
    aba_demais.title = "Demais empresas"
    _escrever_consolidado(
        aba_demais, [linha for linha in linhas if linha.codemp != CODEMP_EMPRESA_SEPARADA]
    )

    aba_separada = livro.create_sheet(_nome_de_aba(_nome_da_empresa_separada(sessao)))
    _escrever_consolidado(
        aba_separada, [linha for linha in linhas if linha.codemp == CODEMP_EMPRESA_SEPARADA]
    )

    buffer = BytesIO()
    livro.save(buffer)
    return buffer.getvalue()
