"""Relatórios de vendas: detalhado e consolidado, sempre com a empresa
separada em aba própria — nomeada com o nome real dela, não "Empresa 17"."""

import uuid
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import delete, select

from app.core import contexto, seguranca
from app.core.db import FabricaDeSessao
from app.excecoes import SemPermissao
from app.models.cadastro import Colaborador, Empresa
from app.models.operacao import Almoco, ItemPedido, Pedido
from app.modules import precos, relatorio_vendas
from app.modules.auditoria import Ator
from tests.conftest import SENHA

PREFIXO = "T-REL-"


@pytest.fixture(autouse=True)
def _contexto_de_teste():
    contexto.definir(correlacao_id=uuid.uuid4(), ip=None, user_agent=None)


@pytest.fixture
def empresa_separada():
    with FabricaDeSessao() as s:
        empresa = s.scalar(
            select(Empresa).where(Empresa.codemp == relatorio_vendas.CODEMP_EMPRESA_SEPARADA)
        )
        assert empresa is not None, "a empresa separada deveria vir da migration de dados iniciais"
        return {"id": empresa.id, "nome": empresa.nome}


@pytest.fixture
def colaborador_empresa_separada(empresa_separada):
    with FabricaDeSessao() as s:
        pessoa = Colaborador(
            nome_completo="Pessoa da Empresa Separada",
            codigo=f"{PREFIXO}SEP",
            codparc=uuid.uuid4().int % 100000 + 800000,
            vinculo="clt",
            matricula=uuid.uuid4().int % 100000 + 800000,
            empresa_id=empresa_separada["id"],
            papel="colaborador",
            senha_hash=seguranca.gerar_hash(SENHA),
        )
        s.add(pessoa)
        s.commit()
        colaborador_id = pessoa.id

    yield colaborador_id

    with FabricaDeSessao() as s:
        s.execute(delete(Pedido).where(Pedido.colaborador_id == colaborador_id))
        s.execute(delete(Almoco).where(Almoco.colaborador_id == colaborador_id))
        s.execute(delete(Colaborador).where(Colaborador.id == colaborador_id))
        s.commit()


def _admin() -> Ator:
    return Ator(id=None, codigo="admin-teste", nome="Admin", papel="admin")


def _pedido_entregue(colaborador_id, empresa_id, *, brinde: bool = False) -> uuid.UUID:
    with FabricaDeSessao() as s:
        colaborador = s.get(Colaborador, colaborador_id)
        pedido = Pedido(
            colaborador_id=colaborador_id,
            empresa_id=empresa_id,
            vinculo=colaborador.vinculo,
            matricula=colaborador.matricula,
            valor_total=Decimal("10.00"),
            status="entregue",
            codigo_retirada=f"{uuid.uuid4().int % 999999:06d}",
        )
        s.add(pedido)
        s.flush()
        s.add(
            ItemPedido(
                pedido_id=pedido.id,
                nome_produto="Refrigerante",
                quantidade=2,
                preco_unitario=Decimal("0.00") if brinde else Decimal("5.00"),
                custo_unitario=Decimal("2.00"),
                brinde=brinde,
            )
        )
        s.commit()
        return pedido.id


def test_detalhado_separa_a_empresa_em_aba_com_nome_real(
    dados, empresa_separada, colaborador_empresa_separada
):
    de = precos.hoje_local()

    _pedido_entregue(dados["ativo"], dados["empresa"])
    _pedido_entregue(colaborador_empresa_separada, empresa_separada["id"])

    with FabricaDeSessao() as s:
        linhas = relatorio_vendas.montar(s, _admin(), de, de)
        livro = load_workbook(BytesIO(relatorio_vendas.planilha_detalhada(s, linhas)))

    nome_esperado = empresa_separada["nome"][:31]
    assert livro.sheetnames == ["Demais empresas", nome_esperado]

    aba_demais = livro["Demais empresas"]
    nomes_demais = [linha[1].value for linha in aba_demais.iter_rows(min_row=2)]
    assert "Fulano de Teste" in nomes_demais
    assert "Pessoa da Empresa Separada" not in nomes_demais

    aba_separada = livro[nome_esperado]
    nomes_separada = [linha[1].value for linha in aba_separada.iter_rows(min_row=2)]
    assert nomes_separada == ["Pessoa da Empresa Separada"]


def test_item_de_brinde_aparece_marcado_na_planilha(dados):
    de = precos.hoje_local()
    _pedido_entregue(dados["ativo"], dados["empresa"], brinde=True)

    with FabricaDeSessao() as s:
        linhas = relatorio_vendas.montar(s, _admin(), de, de)
        assert linhas[0].brinde is True

        livro = load_workbook(BytesIO(relatorio_vendas.planilha_detalhada(s, linhas)))

    cabecalho = [c.value for c in livro["Demais empresas"][1]]
    assert cabecalho[-1] == "Brinde de aniversário"
    linha_da_planilha = next(livro["Demais empresas"].iter_rows(min_row=2))
    assert linha_da_planilha[-1].value == "Sim"


def test_pedido_pendente_nao_entra_no_relatorio(dados):
    de = precos.hoje_local()
    with FabricaDeSessao() as s:
        colaborador = s.get(Colaborador, dados["ativo"])
        pedido = Pedido(
            colaborador_id=dados["ativo"],
            empresa_id=dados["empresa"],
            vinculo=colaborador.vinculo,
            matricula=colaborador.matricula,
            valor_total=Decimal("10.00"),
            status="pendente",
            codigo_retirada=f"{uuid.uuid4().int % 999999:06d}",
        )
        s.add(pedido)
        s.commit()

        linhas = relatorio_vendas.montar(s, _admin(), de, de)
        assert all(linha.colaborador_nome != "Fulano de Teste" for linha in linhas)


def test_colaborador_comum_nao_acessa_o_relatorio(dados):
    comum = Ator(id=dados["ativo"], codigo="x", nome="x", papel="colaborador")
    with FabricaDeSessao() as s:
        with pytest.raises(SemPermissao):
            relatorio_vendas.montar(s, comum, precos.hoje_local(), precos.hoje_local())


def test_consolidado_soma_as_compras_da_mesma_pessoa(dados):
    de = precos.hoje_local()
    _pedido_entregue(dados["ativo"], dados["empresa"])
    _pedido_entregue(dados["ativo"], dados["empresa"])

    with FabricaDeSessao() as s:
        linhas = relatorio_vendas.montar(s, _admin(), de, de)
        consolidado = relatorio_vendas.consolidar_por_pessoa(linhas)

    minha = next(c for c in consolidado if c.colaborador_nome == "Fulano de Teste")
    assert minha.compras == 2
    assert minha.itens == 4
    assert minha.valor_total == Decimal("20.00")


def test_planilha_consolidada_tambem_separa_a_empresa(
    dados, empresa_separada, colaborador_empresa_separada
):
    de = precos.hoje_local()
    _pedido_entregue(dados["ativo"], dados["empresa"])
    _pedido_entregue(colaborador_empresa_separada, empresa_separada["id"])

    with FabricaDeSessao() as s:
        linhas = relatorio_vendas.montar(s, _admin(), de, de)
        consolidado = relatorio_vendas.consolidar_por_pessoa(linhas)
        livro = load_workbook(BytesIO(relatorio_vendas.planilha_consolidada(s, consolidado)))

    nome_esperado = empresa_separada["nome"][:31]
    assert livro.sheetnames == ["Demais empresas", nome_esperado]
